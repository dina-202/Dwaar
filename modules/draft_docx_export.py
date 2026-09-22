"""Deterministic semantic DOCX export for reviewed Dwaar drafts."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from domain.draft_work_product_models import (
    DraftReviewStatus,
    LoadedDraftVersion,
)


def export_draft_docx(
    loaded: LoadedDraftVersion,
    *,
    case_title: str,
) -> bytes:
    """Render reviewed/approved immutable draft text into a DOCX."""
    if not isinstance(loaded, LoadedDraftVersion):
        raise TypeError("loaded must be a LoadedDraftVersion")
    if loaded.metadata.review_status not in {
        DraftReviewStatus.REVIEWED,
        DraftReviewStatus.APPROVED,
    }:
        raise ValueError(
            "only reviewed or approved draft versions may be exported"
        )
    if not isinstance(case_title, str) or not case_title.strip():
        raise ValueError("case_title must be non-empty")
    text = loaded.payload.get("draft_text")
    if not isinstance(text, str) or not text.strip():
        raise ValueError("draft payload text is invalid")

    document = Document()
    document.core_properties.title = case_title.strip()
    document.add_heading(case_title.strip(), level=1)
    document.add_paragraph(
        (
            f"Dwaar draft version {loaded.metadata.version_number} — "
            f"{loaded.metadata.review_status.value.upper()}"
        )
    )

    pending = []
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            if pending:
                document.add_paragraph("\n".join(pending).strip())
                pending = []
            continue
        stripped = line.strip()
        if stripped.startswith("## "):
            if pending:
                document.add_paragraph("\n".join(pending).strip())
                pending = []
            document.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            if pending:
                document.add_paragraph("\n".join(pending).strip())
                pending = []
            document.add_heading(stripped[2:].strip(), level=1)
        else:
            pending.append(line)

    if pending:
        document.add_paragraph("\n".join(pending).strip())

    document.add_paragraph(
        (
            "Source analysis snapshot: "
            f"{loaded.metadata.source_snapshot_id}"
        )
    )
    document.add_paragraph(
        (
            "This document is a professional work-product version. "
            "Its approval state does not record filing."
        )
    )

    output = BytesIO()
    document.save(output)
    return output.getvalue()
